"""
Full backup/restore, Excel readable export, and client ID compaction.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any, Callable

import sqlite3


def _rows(conn: sqlite3.Connection, sql: str, params: tuple = ()) -> list[dict]:
    return conn.execute(sql, params).fetchall()


def build_full_backup_payload(conn: sqlite3.Connection) -> dict:
    """Everything needed for a full restore (except binary document files on disk)."""
    return {
        "backup_format_version": 2,
        "exported_at": datetime.now().isoformat(),
        "clients": _rows(conn, "SELECT * FROM clients WHERE is_deleted = 0 ORDER BY id"),
        "tasks": _rows(conn, "SELECT * FROM tasks WHERE is_deleted = 0 ORDER BY id"),
        "payments": _rows(conn, "SELECT * FROM payments WHERE is_deleted = 0 ORDER BY id"),
        "documents": _rows(conn, "SELECT * FROM documents WHERE is_deleted = 0 ORDER BY id"),
        "client_folder_paths": _rows(conn, "SELECT * FROM client_folder_paths ORDER BY client_id"),
        "deals": _rows(conn, "SELECT * FROM deals ORDER BY id"),
        "deal_partners": _rows(conn, "SELECT * FROM deal_partners ORDER BY id"),
        "deal_partner_receipts": _rows(conn, "SELECT * FROM deal_partner_receipts ORDER BY id"),
        "deal_recruiter_payments": _rows(conn, "SELECT * FROM deal_recruiter_payments ORDER BY id"),
        "deal_transactions": _rows(conn, "SELECT * FROM deal_transactions ORDER BY id"),
        "activities": _rows(conn, "SELECT * FROM activities ORDER BY datetime(created_at) DESC, id DESC LIMIT 500"),
        "settings": {
            r["key"]: r["value"]
            for r in _rows(conn, "SELECT key, value FROM app_settings")
            if not str(r["key"]).startswith("backup_last_")
        },
    }


def wipe_restorable_tables(conn: sqlite3.Connection) -> None:
    conn.executescript(
        """
        DELETE FROM deal_recruiter_payments;
        DELETE FROM deal_partner_receipts;
        DELETE FROM deal_partners;
        DELETE FROM deal_transactions;
        DELETE FROM deals;
        DELETE FROM payments;
        DELETE FROM tasks;
        DELETE FROM documents;
        DELETE FROM client_folder_paths;
        DELETE FROM activities;
        DELETE FROM clients;
        """
    )


def _table_columns(conn: sqlite3.Connection, table: str) -> list[str]:
    return [c["name"] for c in conn.execute(f"PRAGMA table_info({table})").fetchall()]


def _insert_rows(conn: sqlite3.Connection, table: str, rows: list[dict], columns: list[str]) -> None:
    if not rows:
        return
    placeholders = ", ".join(["?"] * len(columns))
    col_sql = ", ".join(columns)
    sql = f"INSERT INTO {table} ({col_sql}) VALUES ({placeholders})"
    for row in rows:
        vals = []
        for c in columns:
            v = row.get(c)
            if v is None and c == "is_deleted" and table in ("clients", "tasks", "payments", "documents"):
                v = 0
            vals.append(v)
        conn.execute(sql, vals)


def _resequence_sqlite_autoincrement(conn: sqlite3.Connection) -> None:
    if not conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name='sqlite_sequence'"
    ).fetchone():
        return
    try:
        for table in (
            "clients",
            "tasks",
            "payments",
            "documents",
            "deals",
            "deal_partners",
            "deal_partner_receipts",
            "deal_recruiter_payments",
            "deal_transactions",
            "activities",
        ):
            try:
                mrow = conn.execute(f"SELECT MAX(id) AS m FROM {table}").fetchone()
            except sqlite3.OperationalError:
                continue
            m = mrow["m"] if mrow else None
            conn.execute("DELETE FROM sqlite_sequence WHERE name = ?", (table,))
            if m:
                conn.execute(
                    "INSERT INTO sqlite_sequence (name, seq) VALUES (?, ?)",
                    (table, int(m)),
                )
    except sqlite3.OperationalError:
        pass


def restore_from_payload(conn: sqlite3.Connection, payload: dict, set_setting: Callable) -> None:
    """
    Replace DB content from backup. v2 = full; v1 = clients/tasks/payments only (deals created on demand).
    """
    wipe_restorable_tables(conn)

    clients = payload.get("clients") or []
    settings = payload.get("settings") or {}

    if not isinstance(clients, list):
        raise ValueError("clients must be a list")

    ver = int(payload.get("backup_format_version") or 1)

    client_id_set = {
        int(c["id"])
        for c in clients
        if isinstance(c, dict) and c.get("id") is not None
    }

    client_cols = _table_columns(conn, "clients")
    _insert_rows(conn, "clients", clients, client_cols)

    tasks = payload.get("tasks") or []
    tasks = [
        t
        for t in tasks
        if isinstance(t, dict)
        and (t.get("client_id") is None or int(t["client_id"]) in client_id_set)
    ]
    payments = payload.get("payments") or []
    payments = [
        p
        for p in payments
        if isinstance(p, dict) and p.get("client_id") is not None and int(p["client_id"]) in client_id_set
    ]
    if tasks:
        _insert_rows(conn, "tasks", tasks, _table_columns(conn, "tasks"))
    if payments:
        _insert_rows(conn, "payments", payments, _table_columns(conn, "payments"))

    if ver >= 2:
        deal_rows = [
            d
            for d in (payload.get("deals") or [])
            if isinstance(d, dict)
            and d.get("client_id") is not None
            and int(d["client_id"]) in client_id_set
        ]
        deal_id_set = {int(d["id"]) for d in deal_rows if d.get("id") is not None}

        def _deal_child(key: str) -> list[dict]:
            return [
                r
                for r in (payload.get(key) or [])
                if isinstance(r, dict)
                and r.get("deal_id") is not None
                and int(r["deal_id"]) in deal_id_set
            ]

        v2_inserts: list[tuple[str, list[dict]]] = [
            ("deals", deal_rows),
            ("deal_partners", _deal_child("deal_partners")),
            ("deal_partner_receipts", _deal_child("deal_partner_receipts")),
            ("deal_recruiter_payments", _deal_child("deal_recruiter_payments")),
            ("deal_transactions", _deal_child("deal_transactions")),
            (
                "documents",
                [
                    r
                    for r in (payload.get("documents") or [])
                    if isinstance(r, dict)
                    and r.get("client_id") is not None
                    and int(r["client_id"]) in client_id_set
                ],
            ),
            (
                "client_folder_paths",
                [
                    r
                    for r in (payload.get("client_folder_paths") or [])
                    if isinstance(r, dict)
                    and r.get("client_id") is not None
                    and int(r["client_id"]) in client_id_set
                ],
            ),
            ("activities", [r for r in (payload.get("activities") or []) if isinstance(r, dict)]),
        ]
        for key, rows in v2_inserts:
            if not rows:
                continue
            cols = _table_columns(conn, key)
            _insert_rows(conn, key, rows, cols)

    for key in ("biz_name", "currency", "backup_dir"):
        if key in settings:
            set_setting(conn, key, str(settings[key]))

    _resequence_sqlite_autoincrement(conn)


def _update_client_fk(conn: sqlite3.Connection, old: int, new: int) -> None:
    conn.execute("UPDATE payments SET client_id = ? WHERE client_id = ?", (new, old))
    conn.execute("UPDATE deals SET client_id = ? WHERE client_id = ?", (new, old))
    conn.execute("UPDATE documents SET client_id = ? WHERE client_id = ?", (new, old))
    conn.execute("UPDATE client_folder_paths SET client_id = ? WHERE client_id = ?", (new, old))
    conn.execute("UPDATE tasks SET client_id = ? WHERE client_id = ?", (new, old))
    conn.execute(
        "UPDATE activities SET ref_id = ? WHERE kind IN ('client', 'finance') AND ref_id = ?",
        (new, old),
    )
    conn.execute(
        "UPDATE activities SET ref_id = ? WHERE kind = 'document' AND ref_id = ? AND message LIKE 'Folder path%'",
        (new, old),
    )


def compact_active_client_ids() -> dict[str, Any]:
    """
    Active clients (is_deleted=0) become sequential #1..n.
    Soft-deleted rows move to high IDs so they never occupy low numbers visible as case #.

    Uses a dedicated autocommit connection so PRAGMA foreign_keys=OFF takes effect.
    (Inside get_db()'s IMMEDIATE transaction, turning FKs off is ignored, and ID remaps fail.)
    """
    from app.database import DB_PATH

    raw = sqlite3.connect(str(DB_PATH), timeout=30, isolation_level=None)
    raw.row_factory = sqlite3.Row
    raw.execute("PRAGMA busy_timeout = 5000")
    try:
        all_rows = raw.execute("SELECT id, is_deleted FROM clients ORDER BY id").fetchall()
        if not all_rows:
            return {"changed": False}

        deleted_ids = [int(r["id"]) for r in all_rows if r["is_deleted"]]
        active_ids = [int(r["id"]) for r in all_rows if not r["is_deleted"]]
        n = len(active_ids)
        desired_seq = list(range(1, n + 1))
        low_blocked = bool(deleted_ids) and min(deleted_ids) <= n
        if active_ids == desired_seq and not low_blocked:
            return {"changed": False}

        max_id = max(int(r["id"]) for r in all_rows)
        temp_high = max(max_id + 1, 200_000) + 25_000

        raw.execute("PRAGMA foreign_keys = OFF")
        try:
            t = temp_high
            for old in sorted(deleted_ids, reverse=True):
                _update_client_fk(raw, old, t)
                raw.execute("UPDATE clients SET id = ? WHERE id = ?", (t, old))
                t += 1

            active_ids = [
                int(r["id"])
                for r in raw.execute(
                    "SELECT id FROM clients WHERE is_deleted = 0 ORDER BY id"
                ).fetchall()
            ]
            if not active_ids:
                _resequence_sqlite_autoincrement(raw)
                return {"changed": True, "active_count": 0}

            mapping = {old: new for old, new in zip(active_ids, range(1, len(active_ids) + 1))}
            if all(mapping[o] == o for o in mapping):
                _resequence_sqlite_autoincrement(raw)
                return {"changed": True, "active_count": len(active_ids)}

            tmp_map = {old: t + i for i, old in enumerate(active_ids)}

            for old, mid in tmp_map.items():
                _update_client_fk(raw, old, mid)
                raw.execute("UPDATE clients SET id = ? WHERE id = ?", (mid, old))

            for old, mid in tmp_map.items():
                final_id = mapping[old]
                _update_client_fk(raw, mid, final_id)
                raw.execute("UPDATE clients SET id = ? WHERE id = ?", (final_id, mid))

            _resequence_sqlite_autoincrement(raw)
            return {"changed": True, "active_count": len(active_ids)}
        finally:
            raw.execute("PRAGMA foreign_keys = ON")

    finally:
        raw.close()


def write_readable_excel(path: Path, conn: sqlite3.Connection) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for Excel backup") from exc

    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="1e3a5f")

    def sheet_from_rows(ws, headers: list[str], rows: list[tuple[Any, ...]]) -> None:
        for col, h in enumerate(headers, 1):
            c = ws.cell(row=1, column=col, value=h)
            c.font = hdr_font
            c.fill = hdr_fill
            c.alignment = Alignment(horizontal="center")
        for r_i, row in enumerate(rows, 2):
            for c_i, val in enumerate(row, 1):
                cell = ws.cell(row=r_i, column=c_i, value=val)
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Clients"
    cheaders = [
        "id",
        "full_name",
        "email",
        "phone",
        "country",
        "visa_type",
        "stage",
        "nationality",
        "passport",
        "process_started_on",
        "total_fee_inr",
        "notes",
        "collected_inr",
        "balance_inr",
        "folder_path",
        "created_at",
        "updated_at",
    ]
    crows = conn.execute(
        """
        SELECT
            c.id,
            c.full_name,
            c.email,
            c.phone,
            c.country,
            c.visa_type,
            c.stage,
            c.nationality,
            c.passport,
            c.process_started_on,
            c.total_fee,
            c.notes,
            COALESCE((
                SELECT SUM(p.amount_paid)
                FROM payments p
                WHERE p.client_id = c.id AND p.is_deleted = 0
            ), 0) + COALESCE((
                SELECT SUM(dpr.amount)
                FROM deals d
                JOIN deal_partner_receipts dpr ON dpr.deal_id = d.id AND dpr.received_from = 'Client'
                WHERE d.client_id = c.id
            ), 0) AS collected_inr,
            (
                c.total_fee - (
                    COALESCE((
                        SELECT SUM(p.amount_paid)
                        FROM payments p
                        WHERE p.client_id = c.id AND p.is_deleted = 0
                    ), 0) + COALESCE((
                        SELECT SUM(dpr.amount)
                        FROM deals d
                        JOIN deal_partner_receipts dpr ON dpr.deal_id = d.id AND dpr.received_from = 'Client'
                        WHERE d.client_id = c.id
                    ), 0)
                )
            ) AS balance_inr,
            (SELECT fp.folder_path FROM client_folder_paths fp WHERE fp.client_id = c.id) AS folder_path,
            c.created_at,
            c.updated_at
        FROM clients c
        WHERE c.is_deleted = 0
        ORDER BY c.id
        """
    ).fetchall()
    sheet_from_rows(
        ws0,
        cheaders,
        [
            (
                r["id"],
                r["full_name"],
                r["email"],
                r["phone"],
                r["country"],
                r["visa_type"],
                r["stage"],
                r["nationality"],
                r["passport"],
                r["process_started_on"],
                r["total_fee"],
                (r["notes"] or "")[:4000],
                r["collected_inr"],
                r["balance_inr"],
                r["folder_path"] or "",
                r["created_at"],
                r["updated_at"],
            )
            for r in crows
        ],
    )

    # Payments (Client → partners)
    wsp = wb.create_sheet(title="Client_to_partner_payments")
    pay_rows = conn.execute(
        """
        SELECT
            p.id,
            p.client_id,
            c.full_name AS client_name,
            p.amount_paid,
            p.payment_type,
            p.collected_by,
            p.description,
            p.paid_at,
            p.created_at
        FROM payments p
        JOIN clients c ON c.id = p.client_id
        WHERE p.is_deleted = 0 AND c.is_deleted = 0
        ORDER BY datetime(COALESCE(p.paid_at, p.created_at)) DESC, p.id DESC
        """
    ).fetchall()
    ph = [
        "payment_id",
        "client_id",
        "client_name",
        "amount_inr",
        "payment_type",
        "collected_by",
        "description",
        "paid_at",
        "created_at",
    ]
    sheet_from_rows(
        wsp,
        ph,
        [
            (
                r["id"],
                r["client_id"],
                r["client_name"],
                r["amount_paid"],
                r["payment_type"],
                r["collected_by"],
                r["description"],
                r["paid_at"],
                r["created_at"],
            )
            for r in pay_rows
        ],
    )

    # Deals (Finance breakdown header fields)
    wsd = wb.create_sheet(title="Deals_finance_summary")
    deal_rows = conn.execute(
        """
        SELECT
            d.id AS deal_id,
            d.client_id,
            c.full_name AS client_name,
            d.deal_value,
            d.declared_costing,
            d.profit_to_split,
            d.profit_split_mode,
            d.recruiting_partners,
            d.recruiter_fees_total,
            d.recruiter_ref_advance,
            d.recruiter_ref_after_permit,
            d.updated_at
        FROM deals d
        JOIN clients c ON c.id = d.client_id AND c.is_deleted = 0
        ORDER BY d.client_id
        """
    ).fetchall()
    dh = [
        "deal_id",
        "client_id",
        "client_name",
        "deal_value",
        "declared_costing",
        "profit_to_split",
        "profit_split_mode",
        "recruiting_partners",
        "recruiter_fees_total",
        "recruiter_ref_advance",
        "recruiter_ref_after_permit",
        "updated_at",
    ]
    sheet_from_rows(
        wsd,
        dh,
        [
            (
                r["deal_id"],
                r["client_id"],
                r["client_name"],
                r["deal_value"],
                r["declared_costing"],
                r["profit_to_split"],
                r["profit_split_mode"],
                r["recruiting_partners"],
                r["recruiter_fees_total"],
                r["recruiter_ref_advance"],
                r["recruiter_ref_after_permit"],
                r["updated_at"],
            )
            for r in deal_rows
        ],
    )

    wspart = wb.create_sheet(title="Deal_partners")
    part_rows = conn.execute(
        """
        SELECT dp.id, dp.deal_id, c.id AS client_id, c.full_name AS client_name, dp.name, dp.share_type, dp.share_value
        FROM deal_partners dp
        JOIN deals d ON d.id = dp.deal_id
        JOIN clients c ON c.id = d.client_id AND c.is_deleted = 0
        ORDER BY c.id, dp.id
        """
    ).fetchall()
    sheet_from_rows(
        wspart,
        ["row_id", "deal_id", "client_id", "client_name", "partner_name", "share_type", "share_value"],
        [
            (r["id"], r["deal_id"], r["client_id"], r["client_name"], r["name"], r["share_type"], r["share_value"])
            for r in part_rows
        ],
    )

    wsrec = wb.create_sheet(title="Money_to_me_receipts")
    rec_rows = conn.execute(
        """
        SELECT dpr.id, d.client_id, c.full_name AS client_name, dpr.amount, dpr.received_in,
               dpr.received_from, dpr.receipt_type, dpr.share_settlement, dpr.received_at, dpr.created_at
        FROM deal_partner_receipts dpr
        JOIN deals d ON d.id = dpr.deal_id
        JOIN clients c ON c.id = d.client_id AND c.is_deleted = 0
        ORDER BY datetime(COALESCE(dpr.received_at, dpr.created_at)) DESC
        """
    ).fetchall()
    sheet_from_rows(
        wsrec,
        [
            "row_id",
            "client_id",
            "client_name",
            "amount_inr",
            "received_in",
            "received_from",
            "receipt_type",
            "share_settlement_inr",
            "received_at",
            "created_at",
        ],
        [
            (
                r["id"],
                r["client_id"],
                r["client_name"],
                r["amount"],
                r["received_in"],
                r["received_from"],
                r["receipt_type"],
                r["share_settlement"],
                r["received_at"],
                r["created_at"],
            )
            for r in rec_rows
        ],
    )

    wsrecr = wb.create_sheet(title="Recruiter_payments")
    rerows = conn.execute(
        """
        SELECT drp.id, c.id AS client_id, c.full_name AS client_name, drp.amount, drp.paid_via, drp.paid_at, drp.created_at
        FROM deal_recruiter_payments drp
        JOIN deals d ON d.id = drp.deal_id
        JOIN clients c ON c.id = d.client_id AND c.is_deleted = 0
        ORDER BY datetime(COALESCE(drp.paid_at, drp.created_at)) DESC
        """
    ).fetchall()
    sheet_from_rows(
        wsrecr,
        ["row_id", "client_id", "client_name", "amount_inr", "paid_via", "paid_at", "created_at"],
        [
            (r["id"], r["client_id"], r["client_name"], r["amount"], r["paid_via"], r["paid_at"], r["created_at"])
            for r in rerows
        ],
    )

    wsdtx = wb.create_sheet(title="Deal_transactions")
    dtx = conn.execute(
        """
        SELECT dt.id, d.client_id, c.full_name AS client_name, dt.txn_type, dt.phase, dt.amount,
               dt.counterparty, dt.payment_mode, dt.notes, dt.txn_at, dt.personal_portion, dt.created_at
        FROM deal_transactions dt
        JOIN deals d ON d.id = dt.deal_id
        JOIN clients c ON c.id = d.client_id AND c.is_deleted = 0
        ORDER BY dt.id
        """
    ).fetchall()
    sheet_from_rows(
        wsdtx,
        [
            "row_id",
            "client_id",
            "client_name",
            "txn_type",
            "phase",
            "amount_inr",
            "counterparty",
            "payment_mode",
            "notes",
            "txn_at",
            "personal_portion",
            "created_at",
        ],
        [
            (
                r["id"],
                r["client_id"],
                r["client_name"],
                r["txn_type"],
                r["phase"],
                r["amount"],
                r["counterparty"],
                r["payment_mode"],
                r["notes"],
                r["txn_at"],
                r["personal_portion"],
                r["created_at"],
            )
            for r in dtx
        ],
    )

    wsdoc = wb.create_sheet(title="Documents_meta")
    docrows = conn.execute(
        """
        SELECT doc.id, doc.client_id, c.full_name AS client_name, doc.original_name, doc.stored_name,
               doc.content_type, doc.file_size, doc.created_at
        FROM documents doc
        JOIN clients c ON c.id = doc.client_id AND c.is_deleted = 0
        WHERE doc.is_deleted = 0
        ORDER BY doc.id
        """
    ).fetchall()
    sheet_from_rows(
        wsdoc,
        [
            "document_id",
            "client_id",
            "client_name",
            "original_name",
            "stored_name_on_server",
            "content_type",
            "file_size_bytes",
            "created_at",
        ],
        [
            (
                r["id"],
                r["client_id"],
                r["client_name"],
                r["original_name"],
                r["stored_name"],
                r["content_type"],
                r["file_size"],
                r["created_at"],
            )
            for r in docrows
        ],
    )

    wsact = wb.create_sheet(title="Activity_log")
    actrows = conn.execute(
        """
        SELECT id, kind, ref_id, message, icon, created_at
        FROM activities
        ORDER BY datetime(created_at) DESC, id DESC
        LIMIT 400
        """
    ).fetchall()
    sheet_from_rows(
        wsact,
        ["id", "kind", "ref_id", "message", "icon", "created_at"],
        [(r["id"], r["kind"], r["ref_id"], r["message"], r["icon"], r["created_at"]) for r in actrows],
    )

    wb.save(path)
